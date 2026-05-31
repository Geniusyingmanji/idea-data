from pathlib import Path
import math
import subprocess
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "materials"
SLIDE_DIR = OUT_DIR / "slides"
W, H = 1920, 1080
FONT = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
FONT_LATIN_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


COLORS = {
    "bg": "#F7F8FA",
    "ink": "#17212B",
    "muted": "#5D6875",
    "line": "#DDE3EA",
    "blue": "#2F6BFF",
    "green": "#1E8E5A",
    "amber": "#C97A14",
    "red": "#C84B4B",
    "teal": "#16858C",
    "violet": "#6652C8",
    "slate": "#344054",
    "white": "#FFFFFF",
}


def hex_to_rgb(value):
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def font(size, bold=False):
    path = FONT_LATIN_BOLD if bold and size <= 30 else FONT
    return ImageFont.truetype(path, size)


def add_wrapped(draw, text, xy, max_width, fnt, fill, line_gap=10):
    x, y = xy
    lines = []
    for para in text.split("\n"):
        if not para:
            lines.append("")
            continue
        current = ""
        for char in para:
            trial = current + char
            if draw.textbbox((0, 0), trial, font=fnt)[2] <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = char
        if current:
            lines.append(current)
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        bbox = draw.textbbox((x, y), line or " ", font=fnt)
        y += (bbox[3] - bbox[1]) + line_gap
    return y


def rect(draw, xy, fill, outline=None, radius=24, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def badge(draw, xy, text, color):
    x, y = xy
    f = font(26)
    pad_x, pad_y = 22, 10
    box = draw.textbbox((0, 0), text, font=f)
    w = box[2] - box[0] + pad_x * 2
    h = box[3] - box[1] + pad_y * 2
    rect(draw, (x, y, x + w, y + h), hex_to_rgb(color), radius=18)
    draw.text((x + pad_x, y + pad_y - 2), text, font=f, fill=hex_to_rgb("#FFFFFF"))
    return w


def header(draw, title, kicker, slide_no):
    draw.text((92, 58), kicker, font=font(26), fill=hex_to_rgb(COLORS["muted"]))
    draw.text((92, 100), title, font=font(54), fill=hex_to_rgb(COLORS["ink"]))
    draw.line((92, 178, 1828, 178), fill=hex_to_rgb(COLORS["line"]), width=2)
    draw.text((1750, 58), f"{slide_no:02d}/10", font=font(26), fill=hex_to_rgb(COLORS["muted"]))


def footer(draw):
    draw.text(
        (92, 1016),
        "SciEvo-Lineage | https://github.com/Geniusyingmanji/idea-data",
        font=font(23),
        fill=hex_to_rgb(COLORS["muted"]),
    )


def card(draw, xy, title, body, accent, title_size=34, body_size=27):
    x1, y1, x2, y2 = xy
    rect(draw, xy, hex_to_rgb(COLORS["white"]), outline=hex_to_rgb(COLORS["line"]), radius=20)
    draw.rectangle((x1, y1, x1 + 10, y2), fill=hex_to_rgb(accent))
    draw.text((x1 + 34, y1 + 28), title, font=font(title_size), fill=hex_to_rgb(COLORS["ink"]))
    add_wrapped(draw, body, (x1 + 34, y1 + 86), x2 - x1 - 68, font(body_size), hex_to_rgb(COLORS["muted"]), 11)


def draw_arrow(draw, start, end, color="#8A95A3", width=5):
    sx, sy = start
    ex, ey = end
    draw.line((sx, sy, ex, ey), fill=hex_to_rgb(color), width=width)
    angle = math.atan2(ey - sy, ex - sx)
    size = 18
    p1 = (ex - size * math.cos(angle - math.pi / 6), ey - size * math.sin(angle - math.pi / 6))
    p2 = (ex - size * math.cos(angle + math.pi / 6), ey - size * math.sin(angle + math.pi / 6))
    draw.polygon((end, p1, p2), fill=hex_to_rgb(color))


def slide_canvas():
    img = Image.new("RGB", (W, H), hex_to_rgb(COLORS["bg"]))
    draw = ImageDraw.Draw(img)
    return img, draw


SLIDES = [
    {
        "title": "SciEvo-Lineage",
        "subtitle": "面向 AI4Science 的 Sci-Evo 类型科研演化数据集",
        "script": "SciEvo-Lineage 是一个面向 AI4Science 的 Sci-Evo 类型数据集。它把单篇论文内的实验闭环和跨论文的领域演化统一成可训练、可评测、可追溯的数据。",
    },
    {
        "title": "为什么需要科学演化数据",
        "script": "传统论文数据集要么只保留摘要和引用，要么只记录单次实验过程。Sci-Evo 任务真正需要的是问题、缺口、决策、验证和失败修正如何随研究推进而变化。",
    },
    {
        "title": "数据集总览",
        "script": "当前 release 共 13050 条结构化记录，覆盖 14 个规范化学科和子领域标签。数据分为 PaperAtom、TransitionAtom、LineageTrajectory 和 AgenticEpisode 四层。",
    },
    {
        "title": "L1 PaperAtom: 单篇科研闭环",
        "script": "第一层严格对齐官方 Sci-Evo 样例，包含初始请求、智能体轨迹、成功验证，并额外加入失败模式。平均每条记录 8.2 个 trajectory step。",
    },
    {
        "title": "L2/L3: 论文间决策与领域演化链",
        "script": "第二层描述相邻论文之间的缺口、假设、决策理由和验证。第三层把多篇论文连成领域演化链，用来学习一个开放问题如何被持续推进。",
    },
    {
        "title": "L4 AgenticEpisode: 科研 Agent 轨迹",
        "script": "第四层把 lineage 上的推理过程转成 ReAct 轨迹，覆盖 search、read、extract、diff 和 propose 等动作，可直接服务于科研 Agent 的 SFT 或偏好训练。",
    },
    {
        "title": "构建流水线",
        "script": "构建流程从 Semantic Scholar、arXiv 和 openAccessPdf 获取公开来源，经 MinerU 或 Docling 解析，再由 GPT-5.5 抽取结构化闭环，最后统一审计和打包。",
    },
    {
        "title": "MinerU 使用方式",
        "script": "本数据集同时使用 MinerU Cloud API 和本地 MinerU2.5-Pro pipeline。所有 L1 记录都保留 parse_tool 字段，便于复现和来源审计。",
    },
    {
        "title": "质量控制与合规",
        "script": "数据构建禁止伪造科学事实。质量控制包括 schema validation、去重、抽样质检和可追溯字段。数据来源限定为公开论文、开放摘要和开放获取 PDF。",
    },
    {
        "title": "提交内容与应用价值",
        "script": "仓库包含数据、样例、原始来源样例、技术报告、schema、审计统计和完整构建代码。它可以用于科学推理训练、科研 Agent、演化 RAG 和 Sci-Evo 风格评测。",
    },
]


def render_slide(idx, spec):
    img, draw = slide_canvas()
    n = idx + 1
    if n == 1:
        draw.rectangle((0, 0, W, H), fill=hex_to_rgb("#EEF4FF"))
        draw.rectangle((0, 0, W, 1080), fill=hex_to_rgb("#F7F8FA"))
        draw.polygon([(1240, 0), (1920, 0), (1920, 1080), (1510, 1080)], fill=hex_to_rgb("#E9F4EF"))
        draw.polygon([(0, 710), (540, 1080), (0, 1080)], fill=hex_to_rgb("#F7E9D7"))
        badge(draw, (92, 90), "Sci-Evo 类型", COLORS["blue"])
        draw.text((92, 185), "SciEvo-Lineage", font=font(86), fill=hex_to_rgb(COLORS["ink"]))
        add_wrapped(draw, spec["subtitle"], (98, 310), 1120, font(42), hex_to_rgb(COLORS["slate"]), 16)
        stats = [("13,050", "结构化记录", COLORS["blue"]), ("4", "层数据结构", COLORS["green"]), ("38", "完整样例", COLORS["amber"]), ("10", "原始来源样例", COLORS["teal"])]
        x = 92
        for value, label, color in stats:
            rect(draw, (x, 520, x + 375, 710), hex_to_rgb(COLORS["white"]), outline=hex_to_rgb("#D7DEE8"), radius=22)
            draw.text((x + 34, 552), value, font=font(56), fill=hex_to_rgb(color))
            draw.text((x + 36, 632), label, font=font(29), fill=hex_to_rgb(COLORS["muted"]))
            x += 414
        draw.text((98, 840), "GitHub: https://github.com/Geniusyingmanji/idea-data", font=font(30), fill=hex_to_rgb(COLORS["ink"]))
        draw.text((98, 894), "数据 CC-BY-4.0 | 代码 MIT | 基于公开论文和开放摘要构建", font=font(29), fill=hex_to_rgb(COLORS["muted"]))
        return img

    header(draw, spec["title"], "SciEvo-Lineage submission material", n)
    footer(draw)

    if n == 2:
        card(draw, (92, 250, 610, 760), "传统摘要/引用数据", "保留标题、摘要、引用边，但缺少实验决策过程。\n难以回答为什么下一篇论文要这样改。", COLORS["amber"])
        card(draw, (700, 250, 1220, 760), "单篇实验轨迹", "能描述一篇论文内部步骤，但缺少跨论文演化关系。\n难以表示领域问题如何持续推进。", COLORS["teal"])
        card(draw, (1310, 250, 1828, 760), "SciEvo-Lineage", "同时记录 intra-paper 闭环和 inter-paper 演化。\n显式建模 gap -> hypothesis -> decision -> validation。", COLORS["blue"])
        draw_arrow(draw, (610, 506), (700, 506))
        draw_arrow(draw, (1220, 506), (1310, 506))
        add_wrapped(draw, "核心目标：让模型学习科学想法如何在真实论文链中被提出、验证、失败和修正。", (190, 835), 1540, font(34), hex_to_rgb(COLORS["ink"]), 14)

    elif n == 3:
        items = [
            ("L1 PaperAtom", "6,010", "单篇论文的科研闭环", COLORS["blue"]),
            ("L2 TransitionAtom", "3,057", "A->B 论文边的决策叙事", COLORS["green"]),
            ("L3 LineageTrajectory", "1,515", "多论文链的领域演化", COLORS["amber"]),
            ("L4 AgenticEpisode", "2,468", "科研 Agent 的 ReAct 轨迹", COLORS["violet"]),
        ]
        xs = [92, 540, 988, 1436]
        for x, (name, count, desc, color) in zip(xs, items):
            rect(draw, (x, 250, x + 390, 715), hex_to_rgb(COLORS["white"]), outline=hex_to_rgb(COLORS["line"]), radius=22)
            draw.rectangle((x, 250, x + 390, 262), fill=hex_to_rgb(color))
            draw.text((x + 32, 300), name, font=font(32), fill=hex_to_rgb(COLORS["ink"]))
            draw.text((x + 32, 390), count, font=font(66), fill=hex_to_rgb(color))
            add_wrapped(draw, desc, (x + 32, 505), 320, font(28), hex_to_rgb(COLORS["muted"]), 12)
        card(draw, (190, 795, 1730, 925), "统一 envelope", "release/data/full_dataset.jsonl.gz 中每行均为 {layer, id, data}，也提供 4 个 layer_*.jsonl.gz 便于按层读取。", COLORS["slate"], 32, 28)

    elif n == 4:
        left = (92, 248, 870, 880)
        rect(draw, left, hex_to_rgb(COLORS["white"]), outline=hex_to_rgb(COLORS["line"]), radius=22)
        rows = [
            ("01_initial_request", "target / input / intent / quantifiable goal"),
            ("02_agent_trajectory", "thought / action / tool / observation / valid"),
            ("03_success_verification", "validation / metrics / final verdict"),
            ("04_failure_modes", "failed step / diagnosis / correction"),
        ]
        y = 292
        for i, (name, body) in enumerate(rows):
            color = [COLORS["blue"], COLORS["green"], COLORS["amber"], COLORS["red"]][i]
            draw.text((132, y), name, font=font(34), fill=hex_to_rgb(color))
            add_wrapped(draw, body, (132, y + 48), 650, font(26), hex_to_rgb(COLORS["muted"]), 9)
            y += 140
        card(draw, (965, 250, 1828, 445), "规模", "6,010 条 L1 PaperAtom；全部通过闭环 schema 校验。", COLORS["blue"], 34, 30)
        card(draw, (965, 492, 1828, 687), "动态过程", "平均 trajectory steps: 8.2；每步强制拆分 Background / Gap / Decision。", COLORS["green"], 34, 30)
        card(draw, (965, 734, 1828, 900), "失败与修正", "平均 failure_modes: 2.96；把失败原因和修正动作作为可学习信号。", COLORS["red"], 34, 30)

    elif n == 5:
        card(draw, (92, 245, 880, 875), "L2 TransitionAtom", "描述相邻论文 A->B 的科研决策：\n\n- gap_in_A: 前作留下什么缺口\n- hypothesis_for_B: 后作提出什么假设\n- decision_rationale: 为什么选择该机制\n- experimental_validation_in_B: 如何验证\n- is_failure_recovery: 是否是失败恢复", COLORS["green"], 36, 28)
        card(draw, (1040, 245, 1828, 875), "L3 LineageTrajectory", "把多篇论文组织成领域演化链：\n\n- open_problem: 起始开放问题\n- chain_trajectory: 每篇里程碑的角色\n- remaining_gap: 每一步留下的后续问题\n- chain_verdict: 当前领域状态", COLORS["amber"], 36, 28)
        draw_arrow(draw, (900, 560), (1020, 560), COLORS["slate"], 6)
        draw.text((842, 505), "论文边", font=font(28), fill=hex_to_rgb(COLORS["muted"]))

    elif n == 6:
        archetypes = [("W3 多论文综合", 1024, COLORS["blue"]), ("W2 单篇扩展", 1021, COLORS["green"]), ("W6 跨域桥接", 423, COLORS["violet"])]
        total = 2468
        x0, y0 = 210, 360
        for i, (label, value, color) in enumerate(archetypes):
            y = y0 + i * 145
            draw.text((x0, y), label, font=font(32), fill=hex_to_rgb(COLORS["ink"]))
            draw.text((x0 + 390, y), str(value), font=font(34), fill=hex_to_rgb(color))
            bar_w = int(760 * value / total)
            rect(draw, (x0 + 520, y + 5, x0 + 1280, y + 46), hex_to_rgb("#E6EBF2"), radius=18)
            rect(draw, (x0 + 520, y + 5, x0 + 520 + bar_w, y + 46), hex_to_rgb(color), radius=18)
        card(draw, (190, 790, 1730, 925), "训练价值", "L4 将 search/read/extract/diff/propose 组织成 ReAct demo，可直接用于科研 Agent 的 SFT、DPO 或工具调用能力评测。", COLORS["violet"], 34, 29)

    elif n == 7:
        steps = [
            ("公开来源", "Semantic Scholar\narXiv / DOI\nopenAccessPdf", COLORS["blue"]),
            ("文档解析", "MinerU Cloud\nMinerU Local\nDocling fallback", COLORS["teal"]),
            ("结构抽取", "GPT-5.5\nSci-Evo schema\nfailure modes", COLORS["green"]),
            ("审计打包", "schema validation\n去重 / 质检\nrelease jsonl.gz", COLORS["amber"]),
        ]
        x = 110
        for i, (title, body, color) in enumerate(steps):
            card(draw, (x, 315, x + 365, 720), title, body, color, 34, 29)
            if i < len(steps) - 1:
                draw_arrow(draw, (x + 372, 520), (x + 435, 520), COLORS["slate"], 5)
            x += 445
        add_wrapped(draw, "所有阶段均保留 provenance：论文 ID、公开外部 ID、解析工具、抽取模型和构建时间戳。", (190, 805), 1540, font(32), hex_to_rgb(COLORS["ink"]), 13)

    elif n == 8:
        data = [("MinerU Cloud", 145, COLORS["blue"]), ("MinerU Local", 1096, COLORS["green"]), ("Docling", 3011, COLORS["amber"]), ("Abstract only", 1758, COLORS["slate"])]
        max_v = max(v for _, v, _ in data)
        x0, y0 = 250, 315
        for i, (label, value, color) in enumerate(data):
            y = y0 + i * 130
            draw.text((x0, y), label, font=font(34), fill=hex_to_rgb(COLORS["ink"]))
            draw.text((x0 + 420, y), f"{value}", font=font(34), fill=hex_to_rgb(color))
            rect(draw, (x0 + 610, y + 4, x0 + 1450, y + 50), hex_to_rgb("#E7ECF3"), radius=19)
            rect(draw, (x0 + 610, y + 4, x0 + 610 + int(840 * value / max_v), y + 50), hex_to_rgb(color), radius=19)
        card(draw, (230, 820, 1690, 935), "复现路径", "Cloud API 用于新增开放 PDF；本地 MinerU2.5-Pro pipeline 用于自托管批处理。每条 L1 记录都写入 parse_tool。", COLORS["teal"], 32, 27)

    elif n == 9:
        boxes = [
            ("真实性", "Prompt 明确禁止 invent facts；实验步骤来自真实论文全文或公开摘要。", COLORS["green"]),
            ("合规", "不访问付费墙内容；不重新分发原始 PDF；保留公开 ID 和 DOI/arXiv 链接。", COLORS["blue"]),
            ("安全", "不含 PII；不含可执行载荷；敏感学科只保留公开方法学层面的结构化描述。", COLORS["red"]),
            ("质量", "Schema 校验、去重、抽样质检、审计统计和数据卡片公开。", COLORS["amber"]),
        ]
        positions = [(92, 255, 900, 500), (1020, 255, 1828, 500), (92, 585, 900, 830), (1020, 585, 1828, 830)]
        for pos, item in zip(positions, boxes):
            card(draw, pos, item[0], item[1], item[2], 36, 29)
        add_wrapped(draw, "Release 校验：L1/L2/L3/L4 valid 均为 100%，完整数据和审计统计随仓库一起提交。", (190, 900), 1540, font(30), hex_to_rgb(COLORS["ink"]), 12)

    elif n == 10:
        card(draw, (92, 250, 910, 520), "提交包", "GitHub 仓库 + 压缩包：\n数据文件、38 条完整样例、10 条原始来源样例、schema、技术报告、构建代码。", COLORS["blue"], 36, 29)
        card(draw, (1010, 250, 1828, 520), "应用场景", "科学推理训练、科研 Agent SFT、演化 RAG、Sci-Evo 风格评测、跨学科方法迁移分析。", COLORS["green"], 36, 29)
        card(draw, (92, 615, 910, 855), "已有影响", "L4 demo 已集成至 idea_train；相关 SFT 相比 baseline 提升约 7 分。", COLORS["amber"], 36, 29)
        card(draw, (1010, 615, 1828, 855), "仓库链接", "https://github.com/Geniusyingmanji/idea-data\n\n类型标注：Sci-Evo", COLORS["violet"], 36, 29)

    return img


def build_images():
    SLIDE_DIR.mkdir(parents=True, exist_ok=True)
    image_paths = []
    for idx, spec in enumerate(SLIDES):
        img = render_slide(idx, spec)
        path = SLIDE_DIR / f"slide_{idx + 1:02d}.png"
        img.save(path, quality=95)
        image_paths.append(path)
    return image_paths


def build_pptx(image_paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in image_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    out = OUT_DIR / "SciEvo-Lineage_Submission.pptx"
    prs.save(out)
    return out


def build_script():
    lines = [
        "# SciEvo-Lineage Overview Video Script",
        "",
        "This script matches `SciEvo-Lineage_Overview.mp4`. The video is a silent slide overview; these notes can be used for live narration or voice-over.",
        "",
    ]
    for idx, spec in enumerate(SLIDES, 1):
        lines.append(f"## Slide {idx}: {spec['title']}")
        lines.append("")
        lines.append(textwrap.fill(spec["script"], width=90))
        lines.append("")
    out = OUT_DIR / "SciEvo-Lineage_Overview_Script.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def build_video(image_paths):
    list_path = OUT_DIR / "video_frames.txt"
    with list_path.open("w", encoding="utf-8") as f:
        for path in image_paths:
            f.write(f"file '{path.resolve()}'\n")
            f.write("duration 7\n")
        f.write(f"file '{image_paths[-1].resolve()}'\n")
    out = OUT_DIR / "SciEvo-Lineage_Overview.mp4"
    cmd = [
        "ffmpeg",
        "-y",
        "-f",
        "concat",
        "-safe",
        "0",
        "-i",
        str(list_path),
        "-f",
        "lavfi",
        "-i",
        "anullsrc=channel_layout=stereo:sample_rate=44100",
        "-shortest",
        "-vf",
        "fps=30,format=yuv420p",
        "-c:v",
        "libx264",
        "-crf",
        "24",
        "-preset",
        "medium",
        "-c:a",
        "aac",
        "-b:a",
        "96k",
        str(out),
    ]
    subprocess.run(cmd, check=True)
    return out


def main():
    image_paths = build_images()
    pptx = build_pptx(image_paths)
    script = build_script()
    video = build_video(image_paths)
    print(f"PPTX: {pptx}")
    print(f"Video: {video}")
    print(f"Script: {script}")
    print(f"Slides: {SLIDE_DIR}")


if __name__ == "__main__":
    main()
